"""Сбор цен по таблицам мониторинга.

Примеры:
    .venv/bin/python run.py                       # все таблицы из config, результат в results/
    .venv/bin/python run.py --tables ТВ --stores wb,ozon
    .venv/bin/python run.py --articles UE32H5000FUXRU,UE32F6000FUXRU
    .venv/bin/python run.py --file ~/Desktop/ТВ.xlsx --out ~/Мониторинг   # любой файл -> ~/Мониторинг/Мониторинг <дата>/
"""
import argparse
import asyncio
import csv
import json
import random
import signal
import sys
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from playwright.async_api import async_playwright

from config import GALAXYSTORE_URL, LINKS_FILE, STORES, TABLES
from excel import as_number, detect_table, write_table
import layout
from make_links_template import link_stores
from screenshot import screenshot
import stores as stores_mod
from stores import STORE_IMPL, Result, StoreError, make_http_client

from paths import DATA_DIR as ROOT, PROFILE_DIR, RESULTS_DIR  # noqa: E402  (рабочие файлы: проект или папка пользователя)
DELAY = {"http": (1, 2.5), "browser": (2, 5)}
JSON_EVENTS = False  # --json: события прогресса строками "@@{...}" для приложения "Мониторинг цен"


class StopWatcher:
    """Остановка прогона: SIGINT/SIGTERM (Mac) или появление стоп-файла (Windows, где сигналов нет)."""

    def __init__(self, stop_file=None):
        self.stop_file = Path(stop_file) if stop_file else None
        self.flag = False
        try:
            loop = asyncio.get_running_loop()
            for sig in (signal.SIGTERM, signal.SIGINT):
                loop.add_signal_handler(sig, self.set)
        except (NotImplementedError, RuntimeError):
            pass  # Windows: только стоп-файл

    def set(self):
        self.flag = True

    @property
    def requested(self):
        return self.flag or bool(self.stop_file and self.stop_file.exists())

    async def watch(self, on_stop):
        """Фоновая задача: как только попросили остановиться — вызвать on_stop()."""
        while not self.requested:
            await asyncio.sleep(1)
        on_stop()


def emit(kind, **data):
    if JSON_EVENTS:
        print("@@" + json.dumps({"type": kind, **data}, ensure_ascii=False), flush=True)


def links_path(args, table_name):
    """Файл ссылок таблицы: --links; с основной папкой (--root) — «Ссылки/links_tv.xlsx» / «links_smart.xlsx»;
    иначе links.xlsx в папке проекта."""
    if args.links:
        return Path(args.links).expanduser()
    return layout.links_file(args.root, table_name, args.legacy_links) if args.root else ROOT / LINKS_FILE


def load_all_links(args, tables):
    links = {}
    for t in tables:
        links.update(load_links(links_path(args, t["name"])))
    return links


def load_links(path):
    """{(лист, артикул): {магазин: ссылка или "-"}}. Колонки — по порядку из config (A артикул, B название, дальше магазины)."""
    if not path.exists():
        print(f"Файл ссылок не найден: {path} — М.Видео/DNS/МТС/samsungstore обновлены не будут")
        return {}
    wb = openpyxl.load_workbook(path, read_only=True)
    links = {}
    for t in TABLES:
        if t["name"] not in wb.sheetnames:
            continue
        stores = link_stores(t)
        for r in wb[t["name"]].iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            cells = {}
            for store, v in zip(stores, r[2:]):
                v = str(v).strip() if v is not None else ""
                if v.startswith("http") or v == "-":
                    cells[store] = v
            links[(t["name"], str(r[0]).strip())] = cells
    return links


def selected_tables(args):
    """Таблицы прогона: выбранный файл (тип определяется по заголовкам) или таблицы из config."""
    if args.file:
        t = detect_table(Path(args.file).expanduser().resolve())
        if not t:
            sys.exit(f"Не удалось определить тип таблицы по заголовкам: {args.file}")
        return [t]
    # колонки — по заголовкам файла (detect_table), а не по буквам из config
    return [detect_table(ROOT / t["file"]) or dict(t, file=str(ROOT / t["file"]))
            for t in TABLES if not args.tables or t["name"] in args.tables]


def output_dir(args):
    """--root: папка дня в основной папке; --out: <папка>/Мониторинг <дата>/; без них — results/ в проекте."""
    if args.root:
        return layout.day_dir(args.root)
    return Path(args.out).expanduser() / f"Мониторинг {date.today().isoformat()}" if args.out else RESULTS_DIR


def build_jobs(args, links, tables):
    """Задания: магазин -> [(таблица, строка, артикул, url, ориентир_цены, модель)]."""
    jobs, missing = {}, []
    for t in tables:
        ws = openpyxl.load_workbook(t["file"]).active
        base_col = t["base_col"]
        n = 0
        for row in range(t["first_row"], ws.max_row + 1):
            article = ws[f"{t['article_col']}{row}"].value
            if not article or not str(article).strip():
                continue
            article = str(article).strip()
            if args.articles and article not in args.articles:
                continue
            n += 1
            if args.limit and n > args.limit:
                break
            row_links = links.get((t["name"], article), {})
            model = str(ws[f"{t['name_col']}{row}"].value or "").strip() if t.get("name_col") else None
            base_old = as_number(ws[f"{base_col}{row}"].value)
            for store, col in t["price_cols"].items():
                if args.stores and store not in args.stores:
                    continue
                _, _, mode = STORE_IMPL[store]
                url = row_links.get(store)
                if mode == "auto":
                    url = GALAXYSTORE_URL.format(article.replace("/", "-"))
                if mode == "link" and not url:
                    missing.append(f"{t['name']}: {article} — {STORES[store]}")
                    continue  # нет ссылки — ячейку не трогаем
                ref = as_number(ws[f"{col}{row}"].value) or base_old
                jobs.setdefault(store, []).append((t["name"], row, article, url, ref, model))
    if missing:
        print(f"Нет ссылки в файле ссылок ({len(missing)}) — эти ячейки не обновятся:")
        print("  " + "\n  ".join(missing))
        print("  Новые артикулы добавляются в файл ссылок командой: .venv/bin/python make_links_template.py")
        emit("missing_links", items=missing)
    return jobs


class Progress:
    """Общий прогресс по всем магазинам: процент, сделано/всего, оценка оставшегося времени.
    Дублируется в заголовок окна Терминала (видно и в свёрнутом окне)."""

    def __init__(self, store_totals):
        self.store_totals = store_totals
        self.total, self.done, self.start, self.eta_sec = sum(store_totals.values()), 0, time.time(), 0
        self.by_store = {}
        try:
            self.tty = open("/dev/tty", "w")
        except OSError:
            self.tty = None

    def tick(self, store, status):
        self.done += 1
        self.by_store.setdefault(store, Counter())[status] += 1
        pct = self.done * 100 // self.total
        # Магазины идут параллельно с разной скоростью — ждать придётся самый медленный
        elapsed = time.time() - self.start
        left = max(elapsed / sum(c.values()) * (self.store_totals[st] - sum(c.values()))
                   for st, c in self.by_store.items())
        self.eta_sec = int(left)
        eta = f"~{int(left // 60)} мин" if left >= 60 else "<1 мин"
        eta = "готово" if self.done == self.total else f"осталось {eta}"
        bar = "█" * (pct // 5) + "░" * (20 - pct // 5)
        if self.tty:
            self.tty.write(f"\033]0;Мониторинг {pct}% · {eta}\007")
            self.tty.flush()
        return f"{bar} {pct:>3}% {self.done}/{self.total} {eta:<17}"

    def retried(self, store, status):
        """Повтор прошёл: ошибка превращается в найденную цену / «нет»."""
        c = self.by_store[store]
        c["ошибка"] -= 1
        c[status] += 1

    def line(self):
        pct = self.done * 100 // self.total
        return f"{'█' * (pct // 5) + '░' * (20 - pct // 5)} {pct:>3}% {self.done}/{self.total} {'повтор':<17}"

    def summary(self):
        lines = ["Итого по магазинам:"]
        for store, c in self.by_store.items():
            lines.append(f"  {STORES[store]:<24} цен: {c['ok']:>3}   нет в наличии: {c['нет']:>3}   ошибок: {c['ошибка']:>3}")
        return "\n".join(lines)


# Повторный проход в конце: ошибки, которые могут пройти со второго раза (антибот, таймаут, сайт не ответил).
# «Ссылка на другой товар», «обновите ссылки» и т.п. повтор не исправит.
RETRY_ROUNDS = 2
RETRY_PAUSE = 30
NO_RETRY = ("другой товар", "обновите ссылки", "не нашёл id", "проверьте вручную")


def retryable(res):
    return not any(m in res.note for m in NO_RETRY)


async def fetch(func, target, url, article, ref, model):
    try:
        return await func(target, url=url, article=article, ref_price=ref, model=model)
    except StoreError as e:
        if "другой товар" in str(e):
            raise
        await asyncio.sleep(15)  # одна повторная попытка: таймауты и антибот обычно проходят
    except Exception:
        await asyncio.sleep(15)
    return await func(target, url=url, article=article, ref_price=ref, model=model)


async def run_store(store, items, page, http, results, log, progress):
    func, needs_browser, _ = STORE_IMPL[store]
    lo, hi = DELAY["browser" if needs_browser else "http"]
    target = page if needs_browser else http
    failed = []

    def report(item, res, status, retry=False):
        tname, row, article, *_ = item
        if retry:
            progress.retried(store, status)
        line = progress.line() if retry else progress.tick(store, status)
        price = res.price if res.price is not None else "-"
        mark = " (повтор)" if retry else ""
        print(f"{line} [{STORES[store]:<24}] {article:<22} {status:<6} {price} {res.note}{mark}", flush=True)
        emit("item", store=STORES[store], table=tname, article=article, status=status, price=res.price, note=res.note,
             url=res.url, done=progress.done, total=progress.total, eta=progress.eta_sec, retry=retry)
        log.append([datetime.now().strftime("%H:%M:%S"), tname, article, STORES[store], status, price, res.url,
                    res.note + mark])

    async def attempt(item):
        tname, row, article, url, ref, model = item
        try:
            res = await fetch(func, target, url, article, ref, model)
            results.setdefault(tname, {})[(row, store)] = res
            return res, "ok" if res.price is not None else "нет"
        except Exception as e:  # любая ошибка магазина не должна ронять весь прогон
            return Result(None, url or "", f"{type(e).__name__}: {e}"), "ошибка"

    for item in items:
        tname, row, article, url, ref, model = item
        if url == "-":  # в links.xlsx отмечено, что товара в магазине нет
            res = Result(None, "", "нет в магазине (по links.xlsx)")
            results.setdefault(tname, {})[(row, store)] = res
            report(item, res, "нет")
            continue
        res, status = await attempt(item)
        report(item, res, status)
        if status == "ошибка" and retryable(res):
            failed.append(item)
        await asyncio.sleep(random.uniform(lo, hi))

    # Артикулы закончились — ещё раз по ошибкам: антибот к этому времени обычно отпускает
    for n in range(1, RETRY_ROUNDS + 1):
        if not failed:
            break
        print(f"\n{STORES[store]}: повторяю ошибки ({len(failed)}), круг {n} из {RETRY_ROUNDS}", flush=True)
        emit("retry_round", store=STORES[store], count=len(failed), round=n, rounds=RETRY_ROUNDS)
        await asyncio.sleep(RETRY_PAUSE)
        still = []
        for item in failed:
            res, status = await attempt(item)
            if status == "ошибка":
                if retryable(res):
                    still.append(item)
            else:
                report(item, res, status, retry=True)
            await asyncio.sleep(random.uniform(lo, hi))
        failed = still


def antibot_hook(watching):
    """Заглушка антибота на странице: окно Chrome — вперёд, сообщение в терминал и приложению."""
    async def hook(page, site, active):
        watching[site] = active
        if active:
            print(f"\n⚠ Антибот на {site}: откройте окно Chrome и пройдите проверку, если она не проходит сама\n",
                  flush=True)
            await bring_to_front(page)
        emit("antibot", site=site, active=active)
    return hook


async def bring_to_front(page):
    try:
        await page.bring_to_front()
        # Свернуть и развернуть — так окно выходит поверх остальных и на Mac, и на Windows
        cdp = await page.context.new_cdp_session(page)
        win = (await cdp.send("Browser.getWindowForTarget"))["windowId"]
        for state in ("minimized", "normal"):
            await cdp.send("Browser.setWindowBounds", {"windowId": win, "bounds": {"windowState": state}})
            await asyncio.sleep(0.3)
        await cdp.detach()
    except Exception:
        pass


async def main(args):
    tables = selected_tables(args)
    stores_mod.ANTIBOT_HOOK = antibot_hook({})
    if not await refresh_mvideo_links(args, tables):
        return
    links = load_all_links(args, tables)
    jobs = build_jobs(args, links, tables)
    if not jobs:
        files = ", ".join(str(links_path(args, t["name"])) for t in tables)
        print("Нечего обрабатывать: проверьте фильтры и ссылки в", files)
        emit("nothing", links=files)
        return
    for store, items in jobs.items():
        print(f"{STORES[store]}: {len(items)} позиций")
    results, log = {}, []
    progress = Progress({k: len(v) for k, v in jobs.items()})
    print(f"Всего: {progress.total} проверок\n")
    emit("plan", total=progress.total, stores=[{"name": STORES[k], "total": len(v)} for k, v in jobs.items()])
    need_browser = any(STORE_IMPL[s][1] for s in jobs)
    async with make_http_client() as http, async_playwright() as pw:
        ctx = None
        if need_browser:
            # Видимое окно: в headless-режиме Wildberries и Ozon не пропускают антибот
            ctx = await pw.chromium.launch_persistent_context(
                str(PROFILE_DIR), channel="chrome", headless=False, locale="ru-RU",
                viewport={"width": 1400, "height": 900},
                args=["--disable-blink-features=AutomationControlled"])
        tasks = []
        for store, items in jobs.items():
            page = await ctx.new_page() if STORE_IMPL[store][1] else None
            tasks.append(asyncio.create_task(run_store(store, items, page, http, results, log, progress)))
        stopped = False

        def stop():
            nonlocal stopped
            stopped = True
            for task in tasks:
                task.cancel()

        watcher = asyncio.create_task(StopWatcher(args.stop_file).watch(stop))
        await asyncio.gather(*tasks, return_exceptions=True)
        watcher.cancel()
        if stopped:
            print(f"\nОстановлено пользователем: обработано {progress.done} из {progress.total}, собранное сохраняется")
            emit("stopped", done=progress.done, total=progress.total)
        if ctx:
            try:
                await ctx.close()
            except Exception:
                pass

    out = output_dir(args)
    out.mkdir(parents=True, exist_ok=True)
    stamp = date.today().isoformat()
    for t in tables:
        if t["name"] not in results:
            continue
        src = Path(t["file"])
        # Основная папка: <день>/ТВ/<исходник>; --out: <папка с датой>/<исходник>; results/ — с датой в имени
        if args.root:
            dst = layout.table_dir(args.root, t["name"]) / src.name
            dst.parent.mkdir(parents=True, exist_ok=True)
        else:
            dst = out / (src.name if args.out else f"{src.stem}_{stamp}{src.suffix}")
        # Повторный запуск в тот же день дописывает в уже собранный файл (например, перепроверка отдельных магазинов)
        changed = write_table(t, dst if dst.exists() else src, dst, results[t["name"]], base_src_path=src)
        print(f"Сохранено: {dst}  (изменено ячеек: {changed})")
        emit("saved", table=t["name"], path=str(dst), changed=changed)
        png = layout.screenshot_path(args.root, t["name"]) if args.root else out / "Скриншоты" / f"{dst.stem}.png"
        if args.root:
            layout.snapshot_links(args.root, t["name"])
        try:
            await screenshot(dst, png)
            print(f"Скриншот: {png}")
            emit("screenshot", table=t["name"], path=str(png))
        except Exception as e:  # скриншот — дополнение, без него результат всё равно сохранён
            print(f"Скриншот не получился: {e}")
    log_dir = layout.logs_dir(args.root) if args.root else out
    log_dir.mkdir(parents=True, exist_ok=True)
    log_path = log_dir / f"Лог {datetime.now():%Y-%m-%d %H-%M-%S}.csv"
    with open(log_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["время", "таблица", "артикул", "магазин", "статус", "цена", "ссылка", "примечание"])
        w.writerows(log)
    print("\n" + progress.summary())
    errors = sum(c["ошибка"] for c in progress.by_store.values())
    print(f"Лог: {log_path}  (ошибок: {errors}; в ячейках с ошибкой оставлены старые значения)")
    emit("summary", errors=errors, log=str(log_path),
         stores={STORES[k]: dict(v) for k, v in progress.by_store.items()})


MVIDEO_REFRESH_LABEL = "М.Видео: обновление ссылок"


async def refresh_mvideo_links(args, tables):
    """Смартфоны: перед сбором цен заново выбираем ссылки М.Видео — у продавцов маркетплейса самое дешёвое
    предложение меняется каждый день (распродано, другой продавец). False — пользователь остановил прогон."""
    phone_tables = [t for t in tables if t.get("autofill") == "phones" and "mvideo" in t["price_cols"]]
    if (args.no_links_refresh or args.articles or args.limit or not phone_tables
            or (args.stores and "mvideo" not in args.stores)):
        return True
    import fill_links  # здесь, а не наверху: fill_links сам импортирует run

    def forward(kind, **data):  # прогресс автозаполнения — отдельной строкой; его итоги («saved» и др.) не нужны
        if kind == "plan":
            emit(kind, total=data["total"], stores=[{"name": MVIDEO_REFRESH_LABEL, "total": data["total"]}])
        elif kind == "item":
            emit(kind, **dict(data, store=MVIDEO_REFRESH_LABEL))

    watcher = StopWatcher(args.stop_file)
    fill_links.run.emit, fill_links.run.StopWatcher = forward, lambda _stop_file: watcher
    for t in phone_tables:
        print(f"Обновляю ссылки М.Видео: {Path(t['file']).name}")
        try:
            await fill_links.fill(argparse.Namespace(
                file=t["file"], links=str(links_path(args, t["name"])), stores="mvideo",
                only_empty=False, only_missing=False, json=False, stop_file=args.stop_file))
        except Exception as e:  # не получилось — собираем цены по прежним ссылкам
            print(f"Ссылки М.Видео не обновлены ({type(e).__name__}: {e}) — цены по прежним ссылкам")
        if watcher.requested:
            print("Остановлено пользователем во время обновления ссылок")
            emit("stopped", done=0, total=0)
            return False
    return True


def parse_args():
    p = argparse.ArgumentParser(description="Мониторинг цен")
    p.add_argument("--tables", type=lambda s: s.split(","), help="ТВ,Смартфоны")
    p.add_argument("--file", help="таблица мониторинга (любой путь); тип определяется по заголовкам")
    p.add_argument("--out", help="папка для результатов; внутри создаётся «Мониторинг <дата>»")
    p.add_argument("--root", help="основная папка: ссылки в «Ссылки», результаты в папке дня (ТВ, Смартфоны, Скриншоты, Логи)")
    p.add_argument("--legacy-links", help="с --root: прежний файл ссылок — переносится в «Ссылки», если там ещё пусто")
    p.add_argument("--links", help=f"файл ссылок (по умолчанию: с --root — «Ссылки/links_tv|links_smart.xlsx», иначе {LINKS_FILE} в папке проекта)")
    p.add_argument("--detect", metavar="FILE", help="только определить тип таблицы и выйти")
    p.add_argument("--stores", type=lambda s: s.split(","), help=",".join(STORES))
    p.add_argument("--articles", type=lambda s: [a.strip() for a in s.split(",")])
    p.add_argument("--limit", type=int, help="не больше N артикулов из каждой таблицы")
    p.add_argument("--no-links-refresh", action="store_true",
                   help="смартфоны: не обновлять ссылки М.Видео перед сбором цен")
    p.add_argument("--json", action="store_true", help="события прогресса для приложения")
    p.add_argument("--stop-file", help="если файл появится — остановиться и сохранить собранное")
    a = p.parse_args()
    unknown = set(a.stores or []) - set(STORES)
    if unknown:
        sys.exit(f"Неизвестные магазины: {unknown}. Доступны: {', '.join(STORES)}")
    return a


if __name__ == "__main__":
    args = parse_args()
    if args.detect:
        t = detect_table(Path(args.detect).expanduser())
        stores = [{"key": k, "name": STORES[k]} for k in t["price_cols"]] if t else []
        print(json.dumps({"table": t and t["name"], "stores": stores}, ensure_ascii=False))
        sys.exit(0 if t else 1)
    JSON_EVENTS = args.json
    asyncio.run(main(args))
