"""Основная папка мониторинга (--root): где что лежит.

    <основная папка>/
        Ссылки/                      рабочие файлы ссылок — правятся здесь, используются каждый день
            links_tv.xlsx
            links_smart.xlsx
            Резервные копии ссылок/
        2026-10-22/                  папка дня
            ТВ/                      таблица ТВ с ценами
            Смартфоны/               таблица смартфонов с ценами
            Ссылки/                  копия файлов ссылок, по которым собирались цены в этот день
            Скриншоты/               ТВ.png, Смартфоны.png
            Логи/
"""
import shutil
from datetime import date
from pathlib import Path

import openpyxl

from config import TABLES
from paths import DATA_DIR

LINKS_DIR = "Ссылки"
LINKS_NAMES = {"ТВ": "links_tv.xlsx", "Смартфоны": "links_smart.xlsx"}
SCREENSHOTS_DIR = "Скриншоты"
LOGS_DIR = "Логи"


def day_dir(root, day=None):
    return Path(root).expanduser() / (day or date.today()).isoformat()


def table_dir(root, table_name):
    return day_dir(root) / table_name


def screenshot_path(root, table_name):
    return day_dir(root) / SCREENSHOTS_DIR / f"{table_name}.png"


def logs_dir(root):
    return day_dir(root) / LOGS_DIR


def links_file(root, table_name, legacy=None):
    """Рабочий файл ссылок таблицы. Если его ещё нет — переносим прежний: legacy (файл, выбранный раньше
    в приложении) или найденный links*.xlsx."""
    path = Path(root).expanduser() / LINKS_DIR / LINKS_NAMES[table_name]
    if not path.exists():
        migrate_links(path, table_name, legacy)
    return path


def snapshot_links(root, table_name):
    """Копия файла ссылок в папку дня: по каким ссылкам собирались цены (повторный запуск в тот же день — обновляет)."""
    src = links_file(root, table_name)
    if src.exists():
        dst = day_dir(root) / LINKS_DIR / src.name
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        return dst


def _link_count(path, sheet):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return 0
    if sheet not in wb.sheetnames:
        return 0
    return sum(1 for row in wb[sheet].iter_rows(min_row=2, values_only=True)
               for v in row[2:] if isinstance(v, str) and v.startswith("http"))


def migrate_links(path, table_name, legacy=None):
    """Первый запуск с основной папкой: переносим файл ссылок, которым пользовались раньше. Кандидаты — файл,
    выбранный в приложении (legacy), и links*.xlsx в известных папках. Берём тот, где на листе этой таблицы
    больше всего ссылок (самый полный); при равенстве — выбранный в приложении, затем самый свежий.
    Листы других таблиц из копии убираем."""
    legacy = Path(legacy).expanduser() if legacy else None
    candidates = [legacy] if legacy and legacy.exists() else []
    for folder in (DATA_DIR, Path.home() / "Documents" / "Мониторинг цен", Path.home() / "Desktop"):
        candidates += [c for c in sorted(Path(folder).glob("links*.xlsx")) if c not in candidates]
    scored = [((n, c == legacy, c.stat().st_mtime), c) for c in candidates if (n := _link_count(c, table_name))]
    if scored:
        _copy_sheet(max(scored, key=lambda x: x[0])[1], path, table_name)


def _copy_sheet(best, path, table_name):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(best)
    for name in [s for s in wb.sheetnames if s != table_name and s in {t["name"] for t in TABLES}]:
        wb.remove(wb[name])
    wb.save(path)
    print(f"Файл ссылок {table_name}: перенесён из {best} в {path}")
