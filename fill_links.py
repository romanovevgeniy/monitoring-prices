"""Автозаполнение файла ссылок (links.xlsx): находит карточки товаров по артикулу (ТВ) или по модели (смартфоны).

    .venv/bin/python fill_links.py --file ~/Desktop/ТВ.xlsx [--links путь/links.xlsx] [--only-empty]

ТВ: М.Видео — по каталогу категорий из config (поиск М.Видео для автоматизированного браузера не работает):
среди карточек с артикулом в названии — самая дешёвая в наличии (любой продавец); в наличии нет — самая дешёвая;
карточек нет — «-». DNS — поиск по артикулу, первая карточка, название которой совпадает с артикулом.
Смартфоны/планшеты/часы/наушники (М.Видео, DNS, МТС, samsungstore) — по модели, памяти, связи: phone_sites.py.
Не нашли — прежняя ссылка остаётся. Ошибки сайта не затирают ссылку. Перед записью — резервная копия файла ссылок.
"""
import argparse
import asyncio
import json
import random
import shutil
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import openpyxl
from playwright.async_api import async_playwright

import run
from config import LINKS_FILE, STORES
from excel import detect_table
from make_links_template import link_stores, table_articles, main as sync_links
from matching import name_matches
from phone_sites import DNS as PhoneDNS, MTS as PhoneMTS, MVideo as PhoneMVideo, SamsungStore as PhoneSamsungStore
from stores import make_http_client, wait_antibot

from paths import DATA_DIR as ROOT  # noqa: E402
SUPPORTED = ("mvideo", "dns")
SUPPORTED_PHONES = ("mvideo", "dns", "mts", "samsungstore")
BROWSER_STORES = ("mvideo", "dns")
DASH = "-"

MV_FETCH = """async body => { const r = await fetch('/bff/products', {method: 'POST', credentials: 'include',
    headers: {'Content-Type': 'application/json'}, body: JSON.stringify(body)}); return [r.status, await r.text()] }"""

DNS_CARDS = """() => [...document.querySelectorAll('a.catalog-product__name, a[class*="product__name"]')]
    .map(a => [a.href.split('?')[0].split('#')[0], (a.innerText || '').trim()])"""


class Found:
    def __init__(self, url, note=""):
        self.url, self.note = url, note


def dns_canonical(url):
    """Для товара не в наличии поиск DNS ведёт на страницу аналогов /product/analog/<id>/… — это та же карточка."""
    return url.replace("/product/analog/", "/product/")


# ---------- М.Видео ----------

async def mvideo_catalog(page, categories):
    """Все карточки Samsung из категорий: [{productId, name, price, status, soldOut}, …].
    Фильтр бренда М.Видео меняет формат (сейчас brand-5-5 / samsung-740095-740095), поэтому id фильтра и значений
    берём из ответа каталога; если фильтра нет — проходим всю категорию и отбираем Samsung по названию."""
    if "mvideo.ru" not in page.url:
        await page.goto("https://www.mvideo.ru/", wait_until="domcontentloaded", timeout=60000)
        await wait_antibot(page)

    async def fetch(body):
        status, text = await page.evaluate(MV_FETCH, body)
        if status != 200:
            raise RuntimeError(f"М.Видео каталог: HTTP {status}")
        return json.loads(text).get("body") or {}

    items = []
    for cat in categories:
        body = {"limit": 60, "cursorId": "", "enrich": True, "sortBy": "popularity", "sortDirection": "desc",
                "filters": [{"id": "category", "valuesId": [cat]}], "isGettingBonusRoubles": True}
        first = await fetch(dict(body, limit=1))
        for f in first.get("filters") or []:
            if str(f.get("id", "")).startswith("brand"):
                values = [o["valueId"] for o in f.get("options") or [] if str(o.get("value", "")).lower().startswith("samsung")]
                if values:
                    body["filters"].append({"id": f["id"], "valuesId": values})
                break
        while True:
            data = await fetch(body)
            items += [x for x in data.get("items") or [] if "samsung" in x.get("name", "").lower()]
            if not data.get("cursorId") or not data.get("items"):
                break
            body["cursorId"] = data["cursorId"]
            await asyncio.sleep(0.5)
    return items


def mvideo_pick(items, article):
    def price(x):
        return (x.get("price") or {}).get("salePrice") or 10 ** 12

    cards = [x for x in items if name_matches(x.get("name", ""), article)]
    if not cards:
        return Found(DASH, "нет в каталоге М.Видео")
    in_stock = [x for x in cards if x.get("status") == "available" and not x.get("soldOut")]
    best = min(in_stock or cards, key=price)
    note = f"{len(cards)} карточек, выбрана самая дешёвая" + ("" if in_stock else " (сейчас нет в наличии)")
    return Found(f"https://www.mvideo.ru/products/{best['productId']}", note)


# ---------- DNS (ТВ) ----------

async def dns_find(page, article):
    await page.goto(f"https://www.dns-shop.ru/search/?q={quote(article)}", wait_until="domcontentloaded", timeout=60000)
    await wait_antibot(page)
    await page.wait_for_timeout(2500)
    if "/product/" in page.url:  # поиск сразу открыл карточку
        name = await page.evaluate("() => (document.querySelector('h1') || {}).innerText || ''")
        if name_matches(name, article):
            return Found(dns_canonical(page.url.split("?")[0].split("#")[0]))
    for href, name in await page.evaluate(DNS_CARDS):
        if "/product/" in href and name_matches(name, article):
            return Found(dns_canonical(href))
    return Found(DASH, "не найдено в DNS")


# ---------- основной цикл ----------

async def fill(args):
    table = detect_table(Path(args.file).expanduser())
    if not table:
        sys.exit(f"Не удалось определить тип таблицы по заголовкам: {args.file}")
    links_path = Path(args.links).expanduser() if args.links else ROOT / LINKS_FILE
    sync_links([table], links_path)  # строки для новых артикулов

    phones = table.get("autofill") == "phones"
    stores = [s for s in link_stores(table) if s in (SUPPORTED_PHONES if phones else SUPPORTED)]
    if "mvideo" in stores and not phones and not table.get("mvideo_categories"):
        stores.remove("mvideo")
    names = {a: n for a, n in table_articles(table)}
    unsupported = [STORES[s] for s in link_stores(table) if s not in stores]
    if unsupported:
        print(f"Автозаполнение пока не умеет: {', '.join(unsupported)} — эти ссылки вставляются вручную")
        run.emit("unsupported", stores=unsupported)
    if not stores:
        run.emit("nothing", links=str(links_path))
        return

    wb = openpyxl.load_workbook(links_path)
    ws = wb[table["name"]]
    col_of = {s: 3 + i for i, s in enumerate(link_stores(table))}
    rows = [(r, str(ws.cell(r, 1).value).strip()) for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
    src_articles = set(names)
    rows = [(r, a) for r, a in rows if a in src_articles]  # только артикулы из выбранной таблицы

    def skip(r, s):
        v = str(ws.cell(r, col_of[s]).value or "").strip()
        if args.only_missing:  # заполняем только пустые и «-»
            return v.startswith("http")
        return args.only_empty and bool(v)

    jobs = {s: [(r, a) for r, a in rows if not skip(r, s)] for s in stores}
    progress = run.Progress({s: len(v) for s, v in jobs.items()})
    print(f"Всего: {progress.total} артикулов × магазинов")
    run.emit("plan", total=progress.total, stores=[{"name": STORES[s], "total": len(v)} for s, v in jobs.items()])
    saved = {"backup": None, "changed": 0}

    def save(store, changes):
        """Записать ссылки магазина сразу, как он закончил (М.Видео быстрый, DNS долгий — не ждём друг друга).
        Перед первой записью — резервная копия файла."""
        changed = [(rc, v) for rc, v in changes.items() if ws.cell(*rc).value != v]
        if not changed:
            return
        if not saved["backup"]:
            backup_dir = links_path.parent / "Резервные копии ссылок"
            backup_dir.mkdir(exist_ok=True)
            saved["backup"] = backup_dir / f"{links_path.stem} {datetime.now():%Y-%m-%d %H-%M-%S}{links_path.suffix}"
            shutil.copy2(links_path, saved["backup"])
            print(f"Резервная копия: {saved['backup']}")
        for (r, c), v in changed:
            ws.cell(r, c).value = v
        wb.save(links_path)
        saved["changed"] += len(changed)
        print(f"{STORES[store]}: записано ссылок {len(changed)} — {links_path}", flush=True)

    def report(store, article, status, value, note):
        # смартфоны ищутся по модели из колонки B — её и показываем, артикул в скобках
        label = f"{names[article].strip()} ({article})" if phones and names.get(article) else article
        line = progress.tick(store, status)
        print(f"{line} [{STORES[store]:<10}] {label:<22} {status:<6} {value or ''} {note}", flush=True)
        run.emit("item", store=STORES[store], table=table["name"], article=label, status=status, price=None,
                 note=(value if value and value != DASH else note) or note, url=value or "",
                 done=progress.done, total=progress.total, eta=progress.eta_sec)

    async def do_store(store, page, http):
        changes = {}
        try:
            await store_loop(store, page, http, changes)
        finally:  # и при остановке — сохраняем то, что успели найти
            save(store, changes)

    def phone_site(store, page, http):
        return {"mvideo": lambda: PhoneMVideo(page, mvideo_catalog), "dns": lambda: PhoneDNS(page),
                "mts": lambda: PhoneMTS(http), "samsungstore": lambda: PhoneSamsungStore(http)}[store]()

    async def store_loop(store, page, http, changes):
        items = None
        site = phone_site(store, page, http) if phones else None
        for row, article in jobs[store]:
            try:
                if site:
                    if store == "mvideo" and site.items is None and page is not None and "mvideo.ru" not in page.url:
                        await page.goto("https://www.mvideo.ru/", wait_until="domcontentloaded", timeout=60000)
                        await wait_antibot(page)
                    url, note = await site.find(names.get(article) or article)
                    found = Found(url or DASH, note)
                    if store == "dns":
                        await asyncio.sleep(random.uniform(1.5, 3.5))
                elif store == "mvideo":
                    if items is None:
                        items = await mvideo_catalog(page, table["mvideo_categories"])
                    found = mvideo_pick(items, article)
                else:
                    found = await dns_find(page, article)
                    await asyncio.sleep(random.uniform(1.5, 3.5))
                current = str(ws.cell(row, col_of[store]).value or "").strip()
                if found.url == DASH and current.startswith("http"):
                    # Не нашли (М.Видео не показывает часть карточек в каталоге) — рабочую ссылку не затираем
                    report(store, article, "ok", current, f"{found.note} — оставлена прежняя ссылка")
                    continue
                changes[(row, col_of[store])] = found.url
                report(store, article, "ok" if found.url != DASH else "нет", found.url, found.note)
            except asyncio.CancelledError:
                raise
            except Exception as e:  # ошибка сайта — прежняя ссылка остаётся
                report(store, article, "ошибка", None, f"{type(e).__name__}: {e}")

    async with async_playwright() as pw, make_http_client() as http:
        ctx = await pw.chromium.launch_persistent_context(
            str(run.PROFILE_DIR), channel="chrome", headless=False, locale="ru-RU",
            viewport={"width": 1400, "height": 900}, args=["--disable-blink-features=AutomationControlled"])
        tasks = [asyncio.create_task(do_store(s, await ctx.new_page() if s in BROWSER_STORES else None, http))
                 for s in stores]
        watcher = asyncio.create_task(run.StopWatcher(args.stop_file).watch(lambda: [t.cancel() for t in tasks]))
        await asyncio.gather(*tasks, return_exceptions=True)
        watcher.cancel()
        try:
            await ctx.close()
        except Exception:
            pass

    print(f"\nОбновлено ссылок: {saved['changed']} — {links_path}")
    print(progress.summary())
    run.emit("saved", table=table["name"], path=str(links_path), changed=saved["changed"])
    errors = sum(c["ошибка"] for c in progress.by_store.values())
    run.emit("summary", errors=errors, stores={STORES[k]: dict(v) for k, v in progress.by_store.items()})


def parse_args():
    p = argparse.ArgumentParser(description="Автозаполнение файла ссылок (ТВ: М.Видео, DNS; смартфоны: + МТС, samsungstore)")
    p.add_argument("--file", required=True, help="таблица мониторинга (тип — по заголовкам)")
    p.add_argument("--links", help=f"файл ссылок (по умолчанию {LINKS_FILE} в папке проекта)")
    p.add_argument("--only-empty", action="store_true", help="заполнять только пустые ячейки")
    p.add_argument("--only-missing", action="store_true", help="заполнять пустые ячейки и ячейки с «-»")
    p.add_argument("--json", action="store_true", help="события прогресса для приложения")
    p.add_argument("--stop-file", help="если файл появится — остановиться и сохранить собранное")
    return p.parse_args()


if __name__ == "__main__":
    a = parse_args()
    run.JSON_EVENTS = a.json
    asyncio.run(fill(a))
