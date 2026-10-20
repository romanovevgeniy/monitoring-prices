"""Excel-лист → HTML, похожий на то, как его показывает Excel. Для скриншотов там, где нет Quick Look (Windows).

Учитывает: ширины колонок и высоты строк (и скрытые колонки), объединённые ячейки, заливку (цвета темы с оттенком,
собственную палитру файла), шрифт (имя, размер, жирность, курсив, цвет), рамки, выравнивание и форматы «#,##0 ₽».
"""
import colorsys
import html
import re
from xml.etree import ElementTree

import openpyxl

DEFAULT_WIDTH = 8.43  # ширина колонки Excel по умолчанию, в символах
DEFAULT_HEIGHT = 15  # высота строки по умолчанию, pt
NBSP = " "
BORDER_CSS = {"thin": "1px solid", "hair": "1px solid", "dotted": "1px dotted", "dashed": "1px dashed",
              "medium": "2px solid", "mediumDashed": "2px dashed", "thick": "3px solid", "double": "3px double"}


def theme_palette(wb):
    """Цвета темы в порядке индексов Excel: lt1, dk1, lt2, dk2, accent1…6, hlink, folHlink."""
    fallback = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47",
                "0563C1", "954F72"]
    xml = getattr(wb, "loaded_theme", None)
    if not xml:
        return fallback
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = ElementTree.fromstring(xml).find(".//a:clrScheme", ns)
    if scheme is None:
        return fallback
    colors = {}
    for el in scheme:
        tag = el.tag.split("}")[1]
        child = el[0]
        colors[tag] = child.get("lastClr") or child.get("val") or "000000"
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
             "hlink", "folHlink"]
    return [colors.get(k, fallback[i]) for i, k in enumerate(order)]


def apply_tint(hex6, tint):
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(v * 255):02X}" for v in (r, g, b))


def css_color(color, palette, default=None, indexed=None):
    if color is None:
        return default
    try:
        if color.type == "rgb" and isinstance(color.rgb, str):
            rgb = color.rgb[-6:]
            return default if color.rgb in ("00000000",) else f"#{rgb}"
        if color.type == "theme" and color.theme is not None and color.theme < len(palette):
            return f"#{apply_tint(palette[color.theme], color.tint or 0)}"
        if color.type == "indexed" and color.indexed is not None:
            from openpyxl.styles.colors import COLOR_INDEX
            table = indexed or COLOR_INDEX  # в файле может быть своя палитра (<indexedColors>)
            if color.indexed < len(table):
                return f"#{table[color.indexed][-6:]}"
    except (TypeError, ValueError):
        pass
    return default


def format_value(value, fmt):
    """Число по формату Excel: разделитель тысяч — пробел, дробная часть — запятая, «₽» если есть в формате."""
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    section = (fmt or "General").split(";")[0]
    if section == "General":
        return f"{value:g}".replace(".", ",")
    decimals = 0
    m = re.search(r"0[.,](0+)", section)
    if m:
        decimals = len(m.group(1))
    text = f"{abs(value):,.{decimals}f}".replace(",", NBSP).replace(".", ",")
    if value < 0:
        text = "-" + text
    if "₽" in section:
        text += NBSP + "₽"
    return text


def sheet_to_html(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    palette = theme_palette(wb)
    indexed = list(wb._colors) if getattr(wb, "_colors", None) else None
    max_row, max_col = ws.max_row, ws.max_column

    merged_origin, merged_hidden = {}, set()
    for rng in ws.merged_cells.ranges:
        merged_origin[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    merged_hidden.add((r, c))

    # openpyxl хранит размеры колонок диапазонами (min..max), а не по одной букве
    widths, hidden = {}, set()
    for dim in ws.column_dimensions.values():
        lo, hi = dim.min or 0, dim.max or 0
        if not lo:
            continue
        for c in range(lo, min(hi, max_col) + 1):
            if dim.width:
                widths[c] = dim.width
            if dim.hidden:
                hidden.add(c)
    cols = [round(widths.get(c, DEFAULT_WIDTH) * 7 + 5) for c in range(1, max_col + 1)]

    out = ['<table style="border-collapse:collapse;table-layout:fixed">', "<colgroup>"]
    out += [f'<col style="width:{w}px">' for c, w in enumerate(cols, 1) if c not in hidden]
    out.append("</colgroup>")
    for r in range(1, max_row + 1):
        dim = ws.row_dimensions.get(r)
        height_pt = dim.height if dim is not None and dim.height else DEFAULT_HEIGHT
        out.append(f'<tr style="height:{round(height_pt * 4 / 3)}px">')
        for c in range(1, max_col + 1):
            if (r, c) in merged_hidden or c in hidden:
                continue
            cell = ws.cell(r, c)
            style = ["overflow:hidden", "padding:0 3px", "white-space:nowrap"]
            if cell.fill is not None and cell.fill.fill_type == "solid":
                bg = css_color(cell.fill.fgColor, palette, indexed=indexed)
                if bg:
                    style.append(f"background:{bg}")
            f = cell.font
            if f is not None:
                style.append(f"font-family:'{f.name or 'Calibri'}',sans-serif")
                style.append(f"font-size:{(f.sz or 11)}pt")
                if f.b:
                    style.append("font-weight:bold")
                if f.i:
                    style.append("font-style:italic")
                color = css_color(f.color, palette, indexed=indexed)
                if color:
                    style.append(f"color:{color}")
            b = cell.border
            for side, name in ((b.left, "left"), (b.right, "right"), (b.top, "top"), (b.bottom, "bottom")):
                if side is not None and side.style:
                    style.append(f"border-{name}:{BORDER_CSS.get(side.style, '1px solid')} "
                                 f"{css_color(side.color, palette, '#000', indexed)}")
            a = cell.alignment
            numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            h = a.horizontal if a is not None and a.horizontal not in (None, "general") else ("right" if numeric else "left")
            style.append(f"text-align:{'center' if h in ('center', 'centerContinuous') else h}")
            v = a.vertical if a is not None and a.vertical else "bottom"
            style.append(f"vertical-align:{'middle' if v == 'center' else v}")
            if a is not None and a.wrap_text:
                style[2] = "white-space:normal"
            span = ""
            if (r, c) in merged_origin:
                rs, cs = merged_origin[(r, c)]
                span = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
            text = html.escape(format_value(cell.value, cell.number_format)).replace("\n", "<br>")
            out.append(f'<td{span} style="{";".join(style)}">{text}</td>')
        out.append("</tr>")
    out.append("</table>")
    return ('<!doctype html><html><head><meta charset="utf-8"></head>'
            '<body style="margin:0;background:#fff">' + "".join(out) + "</body></html>")
