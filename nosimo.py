"""Цены НОСИМО с galaxystore.ru → колонка НОСИМО (base_col) выбранной таблицы.

    .venv/bin/python nosimo.py --file ~/Desktop/ТВ.xlsx [--json]

Заменяет ручное внесение цен из 1С. Цена берётся даже если товара нет в наличии (это наша цена — с ней идёт
сравнение); страницы товара нет (404) — «-»; ошибка сайта — прежнее значение остаётся.
Меняются только значения ячеек колонки НОСИМО; перед записью — копия файла в «Резервные копии таблиц».
"""
import argparse
import asyncio
import random
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

import run
from config import GALAXYSTORE_URL
from excel import detect_table
from stores import ld_offer, make_http_client

STORE_NAME = "galaxystore.ru (НОСИМО)"
DASH = "-"


async def galaxystore_price(client, article):
    url = GALAXYSTORE_URL.format(article.replace("/", "-"))
    r = await client.get(url)
    if r.status_code == 404:
        return DASH, "нет на сайте (404)"
    if r.status_code != 200:
        raise RuntimeError(f"HTTP {r.status_code}")
    offer = ld_offer(r.text)
    if not offer or not offer[0]:
        raise RuntimeError("на странице нет цены")
    price, in_stock, _ = offer
    return price, "" if in_stock else "нет в наличии, цена указана"


async def update(args):
    path = Path(args.file).expanduser()
    table = detect_table(path)
    if not table:
        sys.exit(f"Не удалось определить тип таблицы по заголовкам: {path}")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    col = table["base_col"]
    rows = []
    for r in range(table["first_row"], ws.max_row + 1):
        art = ws[f"{table['article_col']}{r}"].value
        if art and str(art).strip():
            rows.append((r, str(art).strip()))

    progress = run.Progress({"galaxystore": len(rows)})
    print(f"Всего: {len(rows)} артикулов")
    run.emit("plan", total=len(rows), stores=[{"name": STORE_NAME, "total": len(rows)}])
    changes = {}
    stop = run.StopWatcher(args.stop_file)

    async with make_http_client() as client:
        for row, article in rows:
            if stop.requested:
                print("\nОстановлено пользователем — сохраняю то, что успел собрать")
                run.emit("stopped", done=progress.done, total=progress.total)
                break
            try:
                value, note = await galaxystore_price(client, article)
                changes[row] = value
                status = "ok" if value != DASH else "нет"
            except Exception as e:  # ошибка сайта — прежнее значение остаётся
                value, note, status = None, f"{type(e).__name__}: {e}", "ошибка"
            line = progress.tick("galaxystore", status)
            print(f"{line} {article:<22} {status:<6} {value if value is not None else ''} {note}", flush=True)
            run.emit("item", store=STORE_NAME, table=table["name"], article=article, status=status,
                     price=value if isinstance(value, int) else None, note=note, url="",
                     done=progress.done, total=progress.total, eta=progress.eta_sec)
            await asyncio.sleep(random.uniform(0.3, 0.8))

    changed = [(r, v) for r, v in changes.items() if ws[f"{col}{r}"].value != v]
    if changed:
        backup_dir = path.parent / "Резервные копии таблиц"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{path.stem} {datetime.now():%Y-%m-%d %H-%M-%S}{path.suffix}"
        shutil.copy2(path, backup)
        print(f"Резервная копия: {backup}")
        for r, v in changed:
            ws[f"{col}{r}"].value = v
        wb.save(path)
    print(f"\nОбновлено цен НОСИМО: {len(changed)} — {path}")
    run.emit("saved", table=table["name"], path=str(path), changed=len(changed))
    errors = progress.by_store.get("galaxystore", {}).get("ошибка", 0)
    run.emit("summary", errors=errors, stores={STORE_NAME: dict(progress.by_store.get("galaxystore", {}))})


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Цены НОСИМО с galaxystore.ru в колонку НОСИМО таблицы")
    ap.add_argument("--file", required=True)
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--stop-file", help="если файл появится — остановиться и сохранить собранное")
    a = ap.parse_args()
    run.JSON_EVENTS = a.json
    asyncio.run(update(a))
