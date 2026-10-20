"""Запись цен в копию исходной таблицы: меняем только значения и красную подсветку, остальное форматирование не трогаем."""
from collections import Counter
from copy import copy

import openpyxl
from openpyxl.styles import PatternFill

from config import RECOLOR_RED, RED_RGB, RED_THRESHOLD_PCT

DASH = "-"


def is_red(cell):
    return cell.fill is not None and cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb == RED_RGB


def column_default_fills(ws, cols, first_row):
    """Обычная (не красная) заливка каждой колонки — та, что встречается чаще всего."""
    result = {}
    for col in cols:
        counter, samples = Counter(), {}
        for r in range(first_row, ws.max_row + 1):
            c = ws[f"{col}{r}"]
            if is_red(c):
                continue
            key = (c.fill.fill_type, c.fill.fgColor.rgb, c.fill.fgColor.theme, c.fill.fgColor.tint)
            counter[key] += 1
            samples.setdefault(key, c.fill)
        result[col] = copy(samples[counter.most_common(1)[0][0]]) if counter else PatternFill()
    return result


def as_number(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def write_table(table, src_path, dst_path, results, base_src_path=None):
    """results: {(row, store): Result}. Возвращает число изменённых ячеек.
    base_src_path — исходник, откуда берутся цены НОСИМО (из 1С), если src_path — уже собранный ранее результат."""
    wb = openpyxl.load_workbook(src_path)
    ws = wb.active
    base_col = table["base_col"]
    if base_src_path and str(base_src_path) != str(src_path):
        base_ws = openpyxl.load_workbook(base_src_path).active
        for row in range(table["first_row"], ws.max_row + 1):
            ws[f"{base_col}{row}"].value = base_ws[f"{base_col}{row}"].value
    cols = table["price_cols"]
    default_fill = column_default_fills(ws, cols.values(), table["first_row"])
    red_fill = PatternFill("solid", fgColor=RED_RGB)
    changed = 0
    for (row, store), res in results.items():
        cell = ws[f"{cols[store]}{row}"]
        new = res.price if res.price is not None else DASH
        if cell.value != new:
            cell.value = new
            changed += 1
    # Подсветка: красный — у конкурента дешевле, чем у НОСИМО. Пересчитываем все строки,
    # т.к. цены НОСИМО обновляются из 1С независимо от парсера.
    for row in range(table["first_row"], ws.max_row + 1) if RECOLOR_RED else []:
        base = as_number(ws[f"{base_col}{row}"].value)
        for col in cols.values():
            cell = ws[f"{col}{row}"]
            price = as_number(cell.value)
            should_red = bool(base and price is not None and price < base * (1 - RED_THRESHOLD_PCT / 100))
            if should_red and not is_red(cell):
                cell.fill = copy(red_fill)
            elif not should_red and is_red(cell):
                cell.fill = copy(default_fill[col])
    wb.save(dst_path)
    return changed


def detect_table(path):
    """Тип таблицы по заголовкам первой строки и колонки по заголовкам → элемент config.TABLES с file=path, или None."""
    from openpyxl.utils import get_column_letter
    from config import HEADER_WORDS, TABLES
    ws = openpyxl.load_workbook(path, read_only=True).active
    header = {}
    for c in next(ws.iter_rows(max_row=1)):
        if c.value is not None and getattr(c, "column", None):
            header[get_column_letter(c.column)] = str(c.value).lower()
    joined = " ".join(header.values())

    def find(key):
        return next((col for col, text in header.items() if any(w in text for w in HEADER_WORDS[key])), None)

    for t in TABLES:
        if not all(w in joined for w in t["signature"]):
            continue
        found = dict(t, file=str(path))
        found["article_col"] = find("article") or t["article_col"]
        if t.get("name_col"):
            found["name_col"] = find("name") or t["name_col"]
        found["price_cols"] = {s: find(s) or col for s, col in t["price_cols"].items()}
        found["base_col"] = find("base") or t["base_col"]
        return found
    return None
