"""Создаёт или обновляет links.xlsx: по листу на таблицу, строка на артикул, колонка на магазин со ссылками.

Повторный запуск приводит порядок строк к порядку таблицы мониторинга: новые артикулы встают на своё место
(в т.ч. в середину), уже вставленные ссылки сохраняются. Артикулы, которых больше нет в таблицах,
не удаляются — переносятся в конец листа и перечисляются (вдруг вернутся).
"""
import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from config import LINKS_FILE, STORES, TABLES
from stores import STORE_IMPL

from paths import DATA_DIR as ROOT  # noqa: E402


def link_stores(table):
    """Магазины таблицы, для которых нужны ссылки (Ozon/WB ищутся по артикулу)."""
    return [s for s in table["price_cols"] if STORE_IMPL[s][2] == "link"]


def table_articles(table):
    ws = openpyxl.load_workbook(ROOT / table["file"]).active
    out = []
    for r in range(table["first_row"], ws.max_row + 1):
        art = ws[f"{table['article_col']}{r}"].value
        if art and str(art).strip():
            name = ws[f"{table['name_col']}{r}"].value if table["name_col"] else None
            out.append((str(art).strip(), (name or "").strip()))
    return out


def new_sheet(wb, table):
    ws = wb.create_sheet(table["name"])
    ws.append(["Артикул", "Наименование"] + [STORES[s] for s in link_stores(table)])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFDDEBF7")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 28
    for i in range(len(link_stores(table))):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 45
    ws.freeze_panes = "C2"
    return ws


def main(tables=TABLES, path=ROOT / LINKS_FILE):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
    if not path.exists():
        wb.remove(wb.active)
    for t in tables:
        ws = wb[t["name"]] if t["name"] in wb.sheetnames else new_sheet(wb, t)
        width = max(ws.max_column, 2 + len(link_stores(t)))
        existing = {}
        for r in ws.iter_rows(min_row=2, max_col=width, values_only=True):
            if r[0] and str(r[0]).strip():
                existing.setdefault(str(r[0]).strip(), list(r))
        articles = table_articles(t)
        current = {a for a, _ in articles}
        rows, added = [], []
        for a, n in articles:
            if a in existing:
                row = existing[a]
                if n and not row[1]:
                    row[1] = n
            else:
                row = [a, n] + [None] * (width - 2)
                added.append(a)
            rows.append(row)
        gone = [a for a in existing if a not in current]
        rows += [existing[a] for a in gone]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in rows:
            ws.append(row)
        shown = ", ".join(added[:10]) + (f" и ещё {len(added) - 10}" if len(added) > 10 else "")
        print(f"{t['name']}: добавлено {len(added)}" + (f" — {shown}" if added else ""))
        if gone:
            print(f"  нет в таблице мониторинга (перенесены в конец листа): {', '.join(gone)}")
    wb.save(path)
    print(f"Сохранено: {path}. Вставьте ссылки в строки новых артикулов.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Синхронизация links.xlsx с таблицами мониторинга")
    ap.add_argument("--file", help="взять артикулы из этого файла (тип таблицы — по заголовкам)")
    ap.add_argument("--root", help="основная папка: файл ссылок — «Ссылки/links_tv.xlsx» или «links_smart.xlsx»")
    ap.add_argument("--legacy-links", help="с --root: прежний файл ссылок — переносится в «Ссылки», если там ещё пусто")
    ap.add_argument("--links", help=f"файл ссылок (по умолчанию {LINKS_FILE} в папке проекта); если его нет — будет создан")
    a = ap.parse_args()
    links = Path(a.links).expanduser() if a.links else ROOT / LINKS_FILE
    if a.file:
        from excel import detect_table
        t = detect_table(Path(a.file).expanduser().resolve())
        if not t:
            sys.exit(f"Не удалось определить тип таблицы по заголовкам: {a.file}")
        if a.root and not a.links:
            import layout
            links = layout.links_file(a.root, t["name"], a.legacy_links)
        main([t], links)
    elif a.root and not a.links:
        import layout
        for t in TABLES:
            main([t], layout.links_file(a.root, t["name"], a.legacy_links))
    else:
        main(path=links)
